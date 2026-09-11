"""Generate CleanMap one-page summary (Word)."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

OUTPUT = "CleanMap One-Page Summary.docx"

FONT = "Calibri"
TITLE_COLOR = RGBColor(30, 30, 30)
MUTED_COLOR = RGBColor(80, 80, 80)
BODY_SIZE = Pt(9)
HEADING_SIZE = Pt(10)


def shade_cell(cell, fill: str = "E8E8EA") -> None:
    shading = parse_xml(rf'<w:shd {nsdecls("w")} w:fill="{fill}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_margins(cell, top=40, bottom=40, left=80, right=80) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f"</w:tcMar>"
    )
    tcPr.append(tcMar)


def add_run(paragraph, text: str, *, bold: bool = False, size=BODY_SIZE, color=None) -> None:
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = size
    run.font.name = FONT
    if color:
        run.font.color.rgb = color


def add_heading_para(cell, text: str) -> None:
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    add_run(p, text, bold=True, size=HEADING_SIZE, color=TITLE_COLOR)


def add_bullets(cell, items: list[str]) -> None:
    for item in items:
        p = cell.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.12)
        for run in p.runs:
            run.font.size = BODY_SIZE
            run.font.name = FONT


def build() -> Document:
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    # Title block
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    add_run(title, "CleanMap", bold=True, size=Pt(22), color=TITLE_COLOR)

    tagline = doc.add_paragraph()
    tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tagline.paragraph_format.space_after = Pt(6)
    add_run(
        tagline,
        "Salesforce Excel Import Validator & Cleansing Tool  ·  Validate · Cleanse · Import with confidence",
        size=Pt(9),
        color=MUTED_COLOR,
    )

    # Two-column body
    body = doc.add_table(rows=2, cols=2)
    body.style = "Table Grid"
    body.autofit = False
    for row in body.rows:
        for cell in row.cells:
            set_cell_margins(cell)
            shade_cell(cell, "FAFAFA")

    left_top, right_top = body.rows[0].cells
    left_bot, right_bot = body.rows[1].cells

    add_heading_para(left_top, "The Problem")
    add_bullets(
        left_top,
        [
            "Legacy Excel files have inconsistent dates, picklists, and lookups",
            "Failed Salesforce imports cause rework and delayed go-lives",
            "Manual row-by-row review is slow and hard to track",
        ],
    )

    add_heading_para(right_top, "What CleanMap Does")
    add_bullets(
        right_top,
        [
            "Validates & cleanses .xlsx / .xls migration files before import",
            "Auto-detects Salesforce field types from the file (no fixed column list)",
            "Resolves lookup names to Salesforce Ids (optional reference files)",
            "Validates picklist values against object field definitions",
            "Adds Ready / Not Ready status per row; exports for Salesforce Inspector",
        ],
    )

    add_heading_para(left_bot, "How It Works — 6 Steps")
    add_bullets(
        left_bot,
        [
            "Upload main Excel (or load from local path)",
            "Review auto-detected field types",
            "Optionally upload lookup reference files",
            "Optionally upload object field definition for picklists",
            "Select mandatory columns → Run validation",
            "Download cleansed, validated Excel file",
        ],
    )

    add_heading_para(right_bot, "Benefits for the Team")
    add_bullets(
        right_bot,
        [
            "Fewer failed imports and less rework",
            "Hours saved on manual Excel review per batch",
            "Clear readiness view before go-live",
            "Standardized validation across objects and team members",
            "Picklist & lookup issues caught before production",
        ],
    )

    doc.add_paragraph("")

    # Capabilities summary box
    cap_table = doc.add_table(rows=1, cols=1)
    cap_table.style = "Table Grid"
    cap_cell = cap_table.rows[0].cells[0]
    shade_cell(cap_cell, "E8E8EA")
    set_cell_margins(cap_cell, top=60, bottom=60, left=100, right=100)

    cap_title = cap_cell.paragraphs[0]
    cap_title.paragraph_format.space_after = Pt(2)
    add_run(cap_title, "At a Glance — Key Capabilities", bold=True, size=HEADING_SIZE, color=TITLE_COLOR)

    cap_grid = cap_cell.add_table(rows=1, cols=2)
    cap_grid.style = "Table Grid"
    left_cap, right_cap = cap_grid.rows[0].cells
    shade_cell(left_cap, "E8E8EA")
    shade_cell(right_cap, "E8E8EA")
    set_cell_margins(left_cap, top=20, bottom=20, left=60, right=40)
    set_cell_margins(right_cap, top=20, bottom=20, left=40, right=60)

    left_items = [
        "Dynamic schema from 3-row Salesforce headers",
        "Auto-cleanse: trim, blanks, dates, checkboxes",
        "Mandatory column selection",
        "Lookup name → Id resolution",
        "Picklist & multi-select validation",
        "Error cells highlighted in export",
    ]
    right_items = [
        "Ready / Not Ready Status column",
        "Salesforce Inspector-ready export format",
        "Large file support (path or 1 GB upload)",
        "Issues tab + summary-by-column metrics",
        "Supports .xlsx and legacy .xls",
        "Streamlit UI — no coding required",
    ]
    add_bullets(left_cap, left_items)
    add_bullets(right_cap, right_items)

    doc.add_paragraph("")

    # Footer row
    footer = doc.add_table(rows=1, cols=2)
    footer.style = "Table Grid"
    run_cell, link_cell = footer.rows[0].cells
    shade_cell(run_cell, "1C1C1F")
    shade_cell(link_cell, "1C1C1F")
    set_cell_margins(run_cell, top=50, bottom=50, left=90, right=40)
    set_cell_margins(link_cell, top=50, bottom=50, left=40, right=90)

    run_p = run_cell.paragraphs[0]
    add_run(run_p, "Quick Start: ", bold=True, size=Pt(9), color=RGBColor(255, 255, 255))
    add_run(
        run_p,
        "cd excel-validator → source .venv/bin/activate → streamlit run app.py",
        size=Pt(9),
        color=RGBColor(200, 200, 205),
    )

    link_p = link_cell.paragraphs[0]
    link_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(
        link_p,
        "Full docs: CleanMap Documentation.docx  ·  ",
        size=Pt(9),
        color=RGBColor(200, 200, 205),
    )
    add_run(
        link_p,
        "github.com/akhilsitetracker123/CleanMap",
        size=Pt(9),
        color=RGBColor(255, 255, 255),
    )

    return doc


if __name__ == "__main__":
    document = build()
    document.save(OUTPUT)
    print(f"Created: {OUTPUT}")
